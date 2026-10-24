from __future__ import annotations

from pathlib import Path

from docx import Document

from audio_transcriptor.models import TranscriptSegment


def timestamp_srt(seconds: float) -> str:
    milliseconds = int(round(seconds * 1000))
    hours = milliseconds // 3_600_000
    milliseconds %= 3_600_000
    minutes = milliseconds // 60_000
    milliseconds %= 60_000
    secs = milliseconds // 1000
    millis = milliseconds % 1000
    return f"{hours:02}:{minutes:02}:{secs:02},{millis:03}"


def write_txt(path: Path, text: str) -> None:
    path.write_text(text.strip() + "\n", encoding="utf-8")


def write_docx(path: Path, title: str, text: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    for paragraph in text.strip().splitlines():
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())
    document.save(path)


def write_srt(
    path: Path,
    segments: list[TranscriptSegment],
    translated: bool | None = None,
) -> None:
    """Write SRT cues.

    `translated=True` writes the English text only, `translated=False` the source
    text only, and `None` prefers the translation and falls back to the source.
    Empty cues are skipped (with sequential numbering) so an edited transcript
    never ships text the user removed.
    """
    blocks: list[str] = []
    for segment in segments:
        if translated is None:
            text = (segment.translated_text or segment.text).strip()
        elif translated:
            text = segment.translated_text.strip()
        else:
            text = segment.text.strip()
        if not text:
            continue
        blocks.append(
            "\n".join(
                [
                    str(len(blocks) + 1),
                    f"{timestamp_srt(segment.start)} --> {timestamp_srt(segment.end)}",
                    text,
                ]
            )
        )
    path.write_text("\n\n".join(blocks) + "\n", encoding="utf-8")
